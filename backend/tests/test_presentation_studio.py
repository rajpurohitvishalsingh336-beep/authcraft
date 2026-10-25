import io
import uuid
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation as PptxPresentation

from app.main import app

client = TestClient(app)


def get_auth_token():
    """Helper to register and login a test developer."""
    unique_suffix = uuid.uuid4().hex[:8]
    email = f"presentation.tester.{unique_suffix}@example.com"
    password = "SecurePassword123!"

    reg_resp = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": password, "full_name": "Presentation Architect"},
    )
    assert reg_resp.status_code == 201, reg_resp.text

    login_resp = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": password},
    )
    assert login_resp.status_code == 200, login_resp.text
    token = login_resp.json()["access_token"]
    return token, email


def test_presentation_themes_catalog():
    """Verify predefined presentation themes can be queried."""
    resp = client.get("/api/v1/presentations/themes")
    assert resp.status_code == 200
    themes = resp.json()
    assert "modern-indigo" in themes
    assert "emerald-growth" in themes
    assert "slate-corporate" in themes
    assert "minimal-white" in themes


def test_presentation_lifecycle_and_genuine_pptx_export():
    """
    Comprehensive test of Presentation Studio:
    1. AI Deck Generation
    2. Slide Content Editing
    3. New Slide Addition
    4. Slide Deletion
    5. Real PPTX Export parsing and verification via python-pptx
    """
    token, _ = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Generate deck
    gen_payload = {
        "topic": "Zero Trust Authentication for Cloud Native Architectures",
        "audience": "Enterprise CISO & Platform Architects",
        "tone": "Executive",
        "slide_count": 5,
        "template_id": "modern-indigo",
    }
    gen_resp = client.post("/api/v1/presentations/generate", json=gen_payload, headers=headers)
    assert gen_resp.status_code == 200, gen_resp.text
    deck = gen_resp.json()
    assert deck["title"] == "Zero Trust Authentication For Cloud Native Architectures"
    assert deck["slide_count"] >= 5
    presentation_id = deck["id"]
    first_slide = deck["slides"][0]
    second_slide = deck["slides"][1]

    # 2. Update slide
    update_resp = client.put(
        f"/api/v1/presentations/{presentation_id}/slides/{second_slide['id']}",
        json={
            "title": "Updated Threat Landscape Overview",
            "speaker_notes": "Emphasize multi-cloud zero trust policy compliance.",
        },
        headers=headers,
    )
    assert update_resp.status_code == 200
    assert update_resp.json()["title"] == "Updated Threat Landscape Overview"

    # 3. Add a new slide
    add_resp = client.post(
        f"/api/v1/presentations/{presentation_id}/slides",
        json={
            "layout_type": "cards",
            "title": "Compliance Certifications",
            "subtitle": "SOC2 Type II, ISO 27001, and HIPAA Ready",
            "speaker_notes": "Highlight audit trails.",
        },
        headers=headers,
    )
    assert add_resp.status_code == 200
    added_slide_id = add_resp.json()["id"]

    # 4. Fetch presentation details
    detail_resp = client.get(f"/api/v1/presentations/{presentation_id}", headers=headers)
    assert detail_resp.status_code == 200
    detail = detail_resp.json()
    assert detail["slide_count"] == deck["slide_count"] + 1

    # 5. Delete the added slide
    del_resp = client.delete(
        f"/api/v1/presentations/{presentation_id}/slides/{added_slide_id}",
        headers=headers,
    )
    assert del_resp.status_code == 200
    assert del_resp.json()["remaining_count"] == deck["slide_count"]

    # 6. EXPORT TO PPTX & VERIFY WITH REAL PPTX PARSER
    export_resp = client.get(f"/api/v1/presentations/{presentation_id}/export/pptx", headers=headers)
    assert export_resp.status_code == 200
    assert (
        export_resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert len(export_resp.content) > 1000  # Genuine binary file

    # Parse exported binary using python-pptx to prove it is a real PPTX file
    pptx_file = io.BytesIO(export_resp.content)
    parsed_prs = PptxPresentation(pptx_file)

    assert len(parsed_prs.slides) == deck["slide_count"]
    # Check 16:9 widescreen dimensions (13.333 x 7.5 inches = 12192000 x 6858000 EMUs)
    assert round(parsed_prs.slide_width.inches, 2) == 13.33
    assert round(parsed_prs.slide_height.inches, 2) == 7.50

    # Verify slide notes on slide 2
    slide_2 = parsed_prs.slides[1]
    assert slide_2.has_notes_slide
    notes_text = slide_2.notes_slide.notes_text_frame.text
    assert "zero trust" in notes_text.lower()
